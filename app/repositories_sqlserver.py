# app/repositories_sqlserver.py
# (Este archivo habla con la base de datos SQL Server - Simath)
# [CORRECCIÓN FINAL: Protección contra valores NULL que rompen el Dashboard]

import pyodbc
import pandas as pd
import io
import config 
from datetime import datetime, timedelta
from openpyxl.utils import get_column_letter

# --- Traducciones para el Excel ---
PROCESS_TYPE_TRANSLATION = {
    'raw': 'Dato Crudo',
    'pluvio_sum': 'Acumulado Diario',
    'nivel_max': 'Maximo Diario',
    'avg_hourly': 'Promedio por Hora',
    'sum_hourly': 'Acumulado por Hora',
    'max_hourly': 'Maximo por Hora'
}

# ==============================================================================
# === CONEXIÓN A LA BD =========================================================
# ==============================================================================
def get_db_connection(db_key):
    db_config = config.DATABASE_CONNECTIONS.get(db_key)
    if not db_config or db_config['driver'] != 'pyodbc':
        raise ValueError(f"Configuración no válida para SQL Server: {db_key}")

    conn_str = (
        f"DRIVER={{{db_config['odbc_driver']}}};"
        f"SERVER={db_config['host']},{db_config['port']};"
        f"DATABASE={db_config['name']};"
        f"UID={db_config['user']};"
        f"PWD={db_config['pass']};"
    )
    return pyodbc.connect(conn_str)

# ==============================================================================
# === CACHE DE SENSORES ========================================================
# ==============================================================================
def build_active_sensor_cache(db_key):
    print(f"--- Construyendo cache para: {db_key} (SQL Server) ---")
    QUERY = "SELECT DISTINCT idRemotas, idSensores FROM dbo.SensoresRemotas WHERE idRemotas IS NOT NULL;"
    
    temp_cache_db = {}
    try:
        conn = get_db_connection(db_key)
        try:
             df = pd.read_sql_query(QUERY, conn)
             for index, row in df.iterrows():
                 ema_id = int(row['idRemotas'])
                 sensor_id = int(row['idSensores'])
                 if ema_id not in temp_cache_db: temp_cache_db[ema_id] = []
                 temp_cache_db[ema_id].append(sensor_id)
        except Exception as e_pd:
            print(f"Error leyendo cache con pandas: {e_pd}")
        conn.close()
        print(f"¡Cache para {db_key} construido! ({len(temp_cache_db)} EMAs)")
        return temp_cache_db
    except Exception as e:
        print(f"!!! ERROR CRÍTICO cache {db_key}: {e}")
        return {} 

# ==============================================================================
# === HELPER: Coordenadas ======================================================
# ==============================================================================
def dms_to_dd(g, m, s, direccion):
    try:
        val = float(g or 0) + (float(m or 0)/60) + (float(s or 0)/3600)
        return -val if direccion in ['O', 'S'] else val
    except:
        return 0.0

# ==============================================================================
# === REPOSITORIO DE SENSORES ==================================================
# ==============================================================================
def get_sensors_for_ema_repo(db_key, G_SENSOR_CACHE, ema_id):
    db_cache = G_SENSOR_CACHE.get(db_key)
    if not db_cache: return []

    active_ids = set()
    if ema_id == 'todas':
        for ids in db_cache.values(): active_ids.update(ids)
    else:
        active_ids.update(db_cache.get(int(ema_id), []))

    if not active_ids: return []

    try:
        conn = get_db_connection(db_key)
        placeholders = ','.join(['?'] * len(active_ids))
        sql = f"SELECT id, Nombre FROM dbo.Sensores WHERE id IN ({placeholders}) ORDER BY Nombre"
        cursor = conn.cursor()
        cursor.execute(sql, list(active_ids))
        
        sensores = []
        ID_MAP = {7: 'pluviometro', 8: 'bateria', 15: 'presion'}
        
        seen = set()
        for row in cursor.fetchall():
            sid, sname = row
            tipo = ID_MAP.get(sid, 'otro')
            key = sname if ema_id == 'todas' else sid
            if key not in seen:
                seen.add(key)
                sensores.append({
                    'id': sid,
                    'nombre': sname,
                    'table_name': tipo,
                    'search_text': sname.lower()
                })
        
        conn.close()
        return sensores
    except Exception as e:
        print(f"Error get_sensors: {e}")
        return []

# ==============================================================================
# === LÓGICA DE CÁLCULO DE LLUVIA (PYTHON) - CORREGIDO: SUMA DIRECTA ===========
# ==============================================================================
def calcular_lluvia_acumulada(df_raw, agrupar_por='dia'):
    if df_raw.empty: return pd.DataFrame()
    
    # Preparamos las columnas de agrupación
    if agrupar_por == 'dia':
        df_raw['grupo'] = df_raw['tiempo_de_medicion'].dt.date
        col_tiempo = 'dia'
    else:
        df_raw['grupo'] = df_raw['tiempo_de_medicion'].dt.floor('H')
        col_tiempo = 'hora'

    # --- SEPARAR LÓGICA: ARECO vs RESTO ---
    # Asumimos que el nombre contiene 'Areco'
    mask_areco = df_raw['nombre_ema'].str.contains('Areco', case=False, na=False)
    
    df_areco = df_raw[mask_areco].copy()
    df_normal = df_raw[~mask_areco].copy()
    
    resultados = []

    # 1. Lógica para ESTACIONES NORMALES (SUMA DIRECTA)
    # El usuario indica que el sensor envía el dato parcial, no el acumulado histórico.
    # Por lo tanto, SUMAMOS todos los valores del periodo.
    if not df_normal.empty:
        df_agrup_normal = df_normal.groupby(['ema_id', 'nombre_ema', 'grupo'])['valor'].sum().reset_index()
        resultados.append(df_agrup_normal)

    # 2. Lógica para ARECO (MÁXIMO)
    # Areco envía el acumulado ya hecho, buscamos el valor más alto del periodo.
    if not df_areco.empty:
        df_agrup_areco = df_areco.groupby(['ema_id', 'nombre_ema', 'grupo'])['valor'].max().reset_index()
        resultados.append(df_agrup_areco)

    # --- UNIFICAR RESULTADOS ---
    if not resultados:
        return pd.DataFrame()
        
    df_final = pd.concat(resultados, ignore_index=True)
    
    # Formateo final de columnas
    df_final.rename(columns={'grupo': col_tiempo}, inplace=True)
    df_final['descripcion_ema'] = ''
    df_final['sensor_nombre'] = '_Pluviometro'
    df_final['latitud'] = None
    df_final['longitud'] = None
    
    return df_final
# ==============================================================================
# === REPOSITORIO DE REPORTES (SQL Server) =====================================
# ==============================================================================
def generate_report_repo(db_key, ema_id_form, fecha_inicio_str, fecha_fin_str, sensor_info_list, process_type_list):
    
    FECHA_INICIO_SQL = fecha_inicio_str
    FECHA_PREVIA = (datetime.strptime(fecha_inicio_str, '%Y-%m-%d') - timedelta(days=1)).strftime('%Y-%m-%d')
    FECHA_FIN_SQL = (datetime.strptime(fecha_fin_str, '%Y-%m-%d') + timedelta(days=1)).strftime('%Y-%m-%d')
    
    dfs_result = []
    conn = get_db_connection(db_key)
    
    try:
        for sensor_info, process_type in zip(sensor_info_list, process_type_list):
            sensor_id, table_name, sensor_name = sensor_info.split('|')
            sensor_id = int(sensor_id)
            
            # --- CASO ESPECIAL: PLUVIÓMETRO (Acumulados) ---
            if table_name == 'pluviometro' and process_type in ['pluvio_sum', 'sum_hourly']:
                query = f"""
                    SELECT 
                        e.id as ema_id, e.Nombre as nombre_ema, 
                        t.FechaDelDato as tiempo_de_medicion, t.Valor as valor
                    FROM dbo.DatosUTR t
                    JOIN dbo.SensoresRemotas sr ON t.idSensoresRemotas = sr.id
                    JOIN dbo.Remotas e ON sr.idRemotas = e.id
                    WHERE sr.idSensores = 7
                    AND t.FechaDelDato >= ? AND t.FechaDelDato < ?
                """
                params = [FECHA_PREVIA, FECHA_FIN_SQL]
                if ema_id_form != 'todas':
                    query += " AND e.id = ?"
                    params.append(int(ema_id_form))
                
                df_raw = pd.read_sql_query(query, conn, params=params)
                agrupacion = 'dia' if process_type == 'pluvio_sum' else 'hora'
                df_proc = calcular_lluvia_acumulada(df_raw, agrupacion)
                
                if not df_proc.empty:
                    col_filtro = 'dia' if agrupacion == 'dia' else 'hora'
                    df_proc[col_filtro] = pd.to_datetime(df_proc[col_filtro])
                    fecha_inicio_dt = pd.to_datetime(fecha_inicio_str)
                    df_proc = df_proc[df_proc[col_filtro] >= fecha_inicio_dt]
                    df_proc['tipo_procesamiento'] = PROCESS_TYPE_TRANSLATION[process_type]
                    dfs_result.append(df_proc)

            # --- CASO NORMAL ---
            else:
                base_sql = """
                    SELECT e.id AS ema_id, e.Nombre AS nombre_ema, e.Observaciones AS descripcion_ema, 
                    NULL AS latitud, NULL AS longitud, ? AS sensor_nombre,
                """
                cols = ""
                group = ""
                
                if process_type == 'raw':
                    cols = "t.FechaDelDato AS tiempo_de_medicion, NULL AS dia, NULL AS hora, t.Valor AS valor"
                elif process_type == 'nivel_max':
                    cols = "NULL, CAST(t.FechaDelDato as date) as dia, NULL, MAX(t.Valor) as valor"
                    group = "GROUP BY e.id, e.Nombre, e.Observaciones, CAST(t.FechaDelDato as date)"
                else:
                    cols = "t.FechaDelDato AS tiempo_de_medicion, NULL, NULL, t.Valor AS valor"

                sql = f"""
                    {base_sql} {cols}, ? AS tipo_procesamiento
                    FROM dbo.DatosUTR t
                    JOIN dbo.SensoresRemotas sr ON t.idSensoresRemotas = sr.id
                    JOIN dbo.Remotas e ON sr.idRemotas = e.id
                    WHERE sr.idSensores = ?
                    AND t.FechaDelDato >= ? AND t.FechaDelDato < ?
                """
                q_params = [sensor_name, PROCESS_TYPE_TRANSLATION.get(process_type, 'Dato'), sensor_id, FECHA_INICIO_SQL, FECHA_FIN_SQL]
                if ema_id_form != 'todas':
                    sql += " AND e.id = ?"
                    q_params.append(int(ema_id_form))
                if group: sql += f" {group}"
                
                df = pd.read_sql_query(sql, conn, params=q_params)
                dfs_result.append(df)

    except Exception as e:
        print(f"Error generando reporte SQL Server: {e}")
        raise e
    finally:
        conn.close()

    if dfs_result:
        return pd.concat(dfs_result, ignore_index=True)
    else:
        return pd.DataFrame()

# ==============================================================================
# === LISTAS Y MAPAS ===========================================================
# ==============================================================================
def get_ema_list_repo(db_key):
    try:
        conn = get_db_connection(db_key)
        res = pd.read_sql("SELECT id, Nombre FROM dbo.Remotas ORDER BY Nombre", conn).values.tolist()
        conn.close()
        return res
    except: return []

def get_ema_locations_repo(db_key):
    try:
        conn = get_db_connection(db_key)
        cursor = conn.cursor()
        cursor.execute("SELECT id, Nombre, Observaciones, LatGrados, LatMinutos, LatSegundos, LongGrados, LongMinutos, LongSegundos FROM dbo.Remotas WHERE LatGrados IS NOT NULL")
        locs = []
        for r in cursor.fetchall():
            lat = dms_to_dd(r.LatGrados, r.LatMinutos, r.LatSegundos, 'S')
            lon = dms_to_dd(r.LongGrados, r.LongMinutos, r.LongSegundos, 'O')
            if lat != 0 and lon != 0:
                locs.append({'id': r.id, 'nombre': r.Nombre, 'descripcion': r.Observaciones, 'lat': lat, 'lon': lon})
        conn.close()
        return locs
    except: return []

# ==============================================================================
# === DASHBOARD Y POPUP (CORREGIDO: Suma Directa + Lógica Areco) ===============
# ==============================================================================
def get_dashboard_data_repo(db_key, ema_id):
    """
    Trae datos frescos.
    - Si es Areco: El acumulado de hoy es el MAX(Valor) de hoy.
    - Si es Normal: El acumulado de hoy es la SUMA DIRECTA(Valor) de hoy.
    """
    data = {'bateria': None, 'presion': None, 'pluvio_sum_hoy': None}
    
    conn = get_db_connection(db_key)
    try:
        cursor = conn.cursor()

        # 0. IDENTIFICAR SI ES ARECO
        cursor.execute("SELECT Nombre FROM dbo.Remotas WHERE id = ?", (ema_id,))
        row_nombre = cursor.fetchone()
        es_areco = False
        if row_nombre and 'areco' in row_nombre[0].lower():
            es_areco = True

        # 1. Batería (ID 8)
        sql_bat = "SELECT TOP 1 Valor, FechaDelDato FROM dbo.DatosUTR t JOIN dbo.SensoresRemotas sr ON t.idSensoresRemotas=sr.id WHERE sr.idRemotas=? AND sr.idSensores=8 ORDER BY FechaDelDato DESC"
        cursor.execute(sql_bat, (ema_id,))
        row = cursor.fetchone()
        if row and row[0] is not None: 
            data['bateria'] = {'valor': row[0], 'timestamp': row[1]}
        
        # 2. Presión (ID 15)
        sql_pres = "SELECT TOP 1 Valor, FechaDelDato FROM dbo.DatosUTR t JOIN dbo.SensoresRemotas sr ON t.idSensoresRemotas=sr.id WHERE sr.idRemotas=? AND sr.idSensores=15 ORDER BY FechaDelDato DESC"
        cursor.execute(sql_pres, (ema_id,))
        row = cursor.fetchone()
        if row and row[0] is not None: 
            data['presion'] = {'valor': row[0], 'timestamp': row[1]}
        
        # 3. Lluvia (ID 7) - Acumulado HOY
        sql_pluvio = """
            SELECT Valor, FechaDelDato 
            FROM dbo.DatosUTR t 
            JOIN dbo.SensoresRemotas sr ON t.idSensoresRemotas=sr.id 
            WHERE sr.idRemotas=? AND sr.idSensores=7 
            AND FechaDelDato >= CAST(GETDATE() AS date)
            ORDER BY FechaDelDato ASC
        """
        df_pluvio = pd.read_sql_query(sql_pluvio, conn, params=[ema_id])
        
        if not df_pluvio.empty:
            total_hoy = 0.0
            
            if es_areco:
                # LÓGICA ARECO: El acumulado es el valor máximo registrado hoy
                total_hoy = df_pluvio['Valor'].max()
            else:
                # LÓGICA NORMAL (Usuario): Suma directa de todos los valores de hoy
                total_hoy = df_pluvio['Valor'].sum()
            
            if pd.notnull(total_hoy):
                data['pluvio_sum_hoy'] = {'valor': float(total_hoy)}
            
    except Exception as e:
        print(f"Error dashboard SQL: {e}")
    finally:
        conn.close()
        
    return data

# ==============================================================================
# ==============================================================================
def create_excel_from_dataframe(df):
    output = io.BytesIO()
    if 'ema_id' in df.columns: df = df.drop(columns=['ema_id'])
    df.to_excel(output, index=False, sheet_name='Datos', engine='openpyxl')
    return output
